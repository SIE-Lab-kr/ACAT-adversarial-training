"""Export the data behind every paper figure/table to Excel (native charts) and CSV.

Outputs:
  ACAT_figure_data.xlsx            (next to acat_gpu/, one sheet per figure/table, with Excel charts)
  acat_gpu/figure_data/*.csv       (the same tables as plain CSV)
"""
import glob
import json
import os
import pickle

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, ScatterChart, Reference, Series
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
RES = os.path.join(ROOT, "results_gpu")
CSV = os.path.join(ROOT, "figure_data"); os.makedirs(CSV, exist_ok=True)
XLSX = os.path.join(os.path.dirname(ROOT), "ACAT_figure_data.xlsx")

METHODS = ["standard", "pgd_at", "trades", "mart", "acat"]
AT = ["pgd_at", "trades", "mart", "acat"]
MN = {"standard": "Standard", "pgd_at": "PGD-AT", "trades": "TRADES", "mart": "MART", "acat": "ACAT"}
BB = ["vgg11", "resnet18", "wrn28_10"]
BN = {"vgg11": "VGG-11", "resnet18": "ResNet-18", "wrn28_10": "WRN-28-10"}
DN = {"cifar10": "CIFAR-10", "cifar100": "CIFAR-100", "fashionmnist": "Fashion-MNIST", "tinyimagenet": "Tiny-ImageNet (15 cl.)"}
ATT = [("clean", "Clean"), ("fgsm", "FGSM"), ("pgd20", "PGD-20"), ("cw_linf", "CW-Linf"),
       ("autoattack", "AutoAttack"), ("square", "Square (black-box)"), ("transfer", "Transfer (black-box)")]
ABL = [("acat", "ACAT (full)"), ("acat_hardness_only", "Hardness only (lambda=0)"),
       ("acat_boundary_only", "Boundary only (lambda=1)"), ("acat_no_ema", "No EMA (alpha=0)"),
       ("acat_cosine", "Cosine schedule"), ("acat_adv_only", "Adv.-only loss (beta=1)"),
       ("acat_orig_hardfirst", "Hard-first (1-h^n)"), ("acat_orig_hardfirst_adv_only", "Hard-first, beta=1"),
       ("pgd_at", "PGD-AT"), ("trades", "TRADES"), ("mart", "MART")]
COLORS = {"Standard": "8C8C8C", "PGD-AT": "3A73B5", "TRADES": "43A04A", "MART": "8657B0", "ACAT": "D1495B"}


def ev(ds, bb, m):
    p = os.path.join(RES, ds, bb, m, "eval.json")
    return json.load(open(p)) if os.path.exists(p) else None


def hi(ds, bb, m):
    p = os.path.join(RES, ds, bb, m, "history.json")
    return json.load(open(p)) if os.path.exists(p) else None


def acc(e, k):
    v = e.get(k) if e else None
    if v is None:
        return None
    return round(100 * (v["acc"] if isinstance(v, dict) else v), 2)


def mes(rows):
    keys = [k for k, _ in ATT]
    best = {k: max(r[k] for r in rows.values()) for k in keys}
    return {m: round(100 * (1 - np.mean([(best[k] - r[k]) / best[k] for k in keys])), 2) for m, r in rows.items()}


# ============================================================ collect tables
tables = {}   # sheet name -> (dataframe, description, chart spec list)

# --- 1. Main results (Tables III, IV; Figs. 4, 6, 7)
rows = []
for ds in DN:
    for bb in BN:
        evs = {m: ev(ds, bb, m) for m in METHODS}
        if not all(evs[m] for m in AT):
            continue
        at_rows = {m: {k: acc(evs[m], k) for k, _ in ATT} for m in AT}
        mm = mes(at_rows)
        for m in METHODS:
            if not evs[m]:
                continue
            h = hi(ds, bb, m)
            r = {"Dataset": DN[ds], "Backbone": BN[bb], "Method": MN[m]}
            r.update({n: acc(evs[m], k) for k, n in ATT})
            r["MES"] = mm.get(m)
            r["Train time (min)"] = round(h["train_time_sec"] / 60, 1) if h else None
            rows.append(r)
main = pd.DataFrame(rows)
tables["Main_Results"] = (main, "Tables III-IV and Figs. 4, 6, 7: accuracy (%) under every threat model, MES (among the 4 AT methods; Standard excluded) and training time.", [])

# --- 2. Per-attack chart data, one block per setting (Figs. 4 and 6)
pa_rows = []
for (ds, bb), g in main.groupby(["Dataset", "Backbone"], sort=False):
    for _, n in ATT:
        r = {"Setting": f"{ds} / {bb}", "Attack": n}
        for _, m in g.iterrows():
            r[m["Method"]] = m[n]
        pa_rows.append(r)
per_attack = pd.DataFrame(pa_rows)
tables["Fig4_6_PerAttack"] = (per_attack, "Figs. 4 and 6: accuracy (%) per threat model; one block of 7 rows per (dataset, backbone) setting.", [("bar_blocks", "Setting")])

# --- 3. MES (Fig. 7)
mes_df = main[main["Method"] != "Standard"].pivot_table(index=["Dataset", "Backbone"], columns="Method", values="MES", sort=False).reset_index()
mes_df.insert(0, "Setting", mes_df["Dataset"] + " / " + mes_df["Backbone"])
mes_df = mes_df.drop(columns=["Dataset", "Backbone"])[["Setting", "PGD-AT", "TRADES", "MART", "ACAT"]]
tables["Fig7_MES"] = (mes_df, "Fig. 7: MES (%) of the four AT methods per setting.", [("bar", "MES (%)")])

# --- 4. Trade-off scatter (Fig. 5)
sc = main[(main["Dataset"] == "CIFAR-10") & (main["Method"] != "Standard")][["Backbone", "Method", "Clean", "AutoAttack"]]
tables["Fig5_Tradeoff"] = (sc.reset_index(drop=True), "Fig. 5: clean vs. AutoAttack accuracy (%) on CIFAR-10, AT methods only.", [("scatter_methods", None)])

# --- 5. eps curve and PGD-steps curve (Fig. 8)
for key, sheet, xname, desc in [("eps_curve", "Fig8a_Eps_Curve", "epsilon (x/255)", "Fig. 8a: PGD-20 robust accuracy (%) vs. perturbation budget, first 1,000 CIFAR-10 test images."),
                                ("steps_curve", "Fig8b_PGD_Steps", "PGD iterations", "Fig. 8b: robust accuracy (%) at eps=8/255 vs. number of PGD iterations, first 1,000 CIFAR-10 test images.")]:
    blocks = []
    for bb in BB:
        evs = {m: ev("cifar10", bb, m) for m in METHODS}
        ks = sorted(evs["acat"][key], key=float)
        for k in ks:
            r = {"Backbone": BN[bb], xname: float(k)}
            r.update({MN[m]: round(100 * evs[m][key][k], 2) for m in METHODS})
            blocks.append(r)
    tables[sheet] = (pd.DataFrame(blocks), desc, [("line_blocks", "Backbone", xname)])

# --- 6. Training dynamics (Fig. 10): per backbone, per epoch, per method
td = []
for bb in BB:
    H = {m: hi("cifar10", bb, m) for m in METHODS}
    for i in range(len(H["acat"]["epochs"])):
        r = {"Backbone": BN[bb], "Epoch": i + 1}
        for m in METHODS:
            e = H[m]["epochs"][i]
            r[f"{MN[m]} loss"] = round(e["train_loss"], 4)
            r[f"{MN[m]} clean acc"] = round(100 * e["clean_acc"], 2)
            r[f"{MN[m]} PGD-10 acc"] = round(100 * e["pgd10_acc"], 2)
        td.append(r)
tables["Fig10_Training_Dynamics"] = (pd.DataFrame(td), "Fig. 10: per-epoch training loss, clean accuracy (%) and PGD-10 accuracy (%) on the first 1,000 CIFAR-10 test images.", [("line_dyn", None)])

# all other datasets' training curves too (not plotted in the paper, but saved)
td2 = []
for ds in ["cifar100", "fashionmnist", "tinyimagenet"]:
    H = {m: hi(ds, "resnet18", m) for m in METHODS}
    if not all(H.values()):
        continue
    for i in range(len(H["acat"]["epochs"])):
        r = {"Dataset": DN[ds], "Epoch": i + 1}
        for m in METHODS:
            e = H[m]["epochs"][i]
            r[f"{MN[m]} loss"] = round(e["train_loss"], 4)
            r[f"{MN[m]} clean acc"] = round(100 * e["clean_acc"], 2)
            r[f"{MN[m]} PGD-10 acc"] = round(100 * e["pgd10_acc"], 2)
        td2.append(r)
tables["Extra_Dynamics_OtherData"] = (pd.DataFrame(td2), "Supplementary (no paper figure): per-epoch curves for ResNet-18 on CIFAR-100, Fashion-MNIST and Tiny-ImageNet.", [("line_dyn_ds", None)])

# --- 7. Curriculum dynamics (Fig. 11)
cd = []
for bb in BB:
    for e in hi("cifar10", bb, "acat")["epochs"]:
        cd.append({"Backbone": BN[bb], "Epoch": e["epoch"] + 1, "lambda_t": round(e["lambda"], 4),
                   "corr(w,h)": round(e["corr_w_h"], 4), "corr(w,b_smooth)": round(e["corr_w_b"], 4),
                   "std(w)": round(e["w_std"], 4), "w 10th pct": round(e["w_p10"], 4), "w 90th pct": round(e["w_p90"], 4),
                   "mean h": round(e["h_mean"], 4), "mean b_smooth": round(e["b_smooth_mean"], 4), "mean b (raw)": round(e["b_mean"], 4)})
tables["Fig11_Curriculum"] = (pd.DataFrame(cd), "Fig. 11: internal ACAT curriculum statistics over the whole training set, per epoch (CIFAR-10).", [("line_cur", None)])

# --- 8. Weights vs hardness (Fig. 12): binned means (line chart) + raw subsample
files = sorted(glob.glob(os.path.join(RES, "cifar10", "resnet18", "acat", "acat_signals_ep*.npz")), key=lambda f: int(f.split("ep")[-1].split(".")[0]))
bins = np.linspace(-1, 1, 21)
binned, raw = [], []
for f, stage in zip(files, ["early", "middle", "final"]):
    d = np.load(f)["data"]; ep = int(f.split("ep")[-1].split(".")[0]) + 1
    idx = np.digitize(d[:, 1], bins) - 1
    for b in range(20):
        sel = idx == b
        if stage == "early":
            binned.append({"h bin center": round((bins[b] + bins[b + 1]) / 2, 2)})
        binned[b][f"mean w, {stage} (ep {ep})"] = round(float(d[sel, 4].mean()), 4) if sel.any() else None
        binned[b][f"n, {stage}"] = int(sel.sum())
    sub = np.random.default_rng(0).choice(len(d), 2000, replace=False)
    for i in sub:
        raw.append({"Stage": stage, "Epoch": ep, "sample index": int(d[i, 0]), "h": round(float(d[i, 1]), 4),
                    "b (raw)": round(float(d[i, 2]), 4), "b_smooth": round(float(d[i, 3]), 4), "w": round(float(d[i, 4]), 4)})
tables["Fig12_Weights_vs_h"] = (pd.DataFrame(binned), "Fig. 12: mean curriculum weight w per hardness bin at the early, middle and final epoch (CIFAR-10, ResNet-18, all 49,920 training samples per epoch).", [("line_bins", None)])
tables["Fig12_Raw_Sample"] = (pd.DataFrame(raw), "Fig. 12 raw points: random 2,000 training samples per stage (h, b, smoothed b, w).", [])

# --- 9. Per-class PGD-20 (Fig. 9) for all three backbones
with open(os.path.join(os.path.expanduser("~"), ".cache", "torch", "cifar-10-batches-py", "test_batch"), "rb") as fh:
    y = np.array(pickle.load(fh, encoding="latin1")["labels"])
CLS = ["plane", "car", "bird", "cat", "deer", "dog", "frog", "horse", "ship", "truck"]
pc = []
for bb in BB:
    for m in METHODS:
        pr = np.array(ev("cifar10", bb, m)["pgd20"]["preds"])
        cl = np.array(ev("cifar10", bb, m)["clean"]["preds"])
        r = {"Backbone": BN[bb], "Method": MN[m], "Metric": "PGD-20 acc"}
        r.update({c: round(100 * float(np.mean(pr[y == i] == i)), 2) for i, c in enumerate(CLS)})
        pc.append(r)
        r = {"Backbone": BN[bb], "Method": MN[m], "Metric": "Clean acc"}
        r.update({c: round(100 * float(np.mean(cl[y == i] == i)), 2) for i, c in enumerate(CLS)})
        pc.append(r)
tables["Fig9_PerClass"] = (pd.DataFrame(pc), "Fig. 9: per-class clean and PGD-20 accuracy (%) on the full CIFAR-10 test set (paper shows ResNet-18 PGD-20).", [("bar_perclass", None)])

# --- 10. Ablation (Table V, Fig. 13)
ab = []
for k, n in ABL:
    e = ev("cifar10", "resnet18", k); h = hi("cifar10", "resnet18", k)
    r = {"Variant": n, "folder": k, "beta": 1.0 if "adv_only" in k else (0.5 if k.startswith("acat") else None)}
    r.update({nn: acc(e, kk) for kk, nn in ATT})
    r["Train time (min)"] = round(h["train_time_sec"] / 60, 1)
    ab.append(r)
tables["Tab5_Fig13_Ablation"] = (pd.DataFrame(ab), "Table V and Fig. 13: ACAT ablations and baselines, CIFAR-10 / ResNet-18 (accuracy %).", [("scatter_ablation", None)])

# --- 11. Cost (Table VI)
co = main[main["Method"] != "Standard"].pivot_table(index=["Dataset", "Backbone"], columns="Method", values="Train time (min)", sort=False).reset_index()
co["COR ACAT"] = (co["ACAT"] / co["PGD-AT"]).round(3)
co["COR TRADES"] = (co["TRADES"] / co["PGD-AT"]).round(3)
co["COR MART"] = (co["MART"] / co["PGD-AT"]).round(3)
co.insert(0, "Setting", co["Dataset"] + " / " + co["Backbone"])
co = co.drop(columns=["Dataset", "Backbone"])[["Setting", "PGD-AT", "TRADES", "MART", "ACAT", "COR TRADES", "COR MART", "COR ACAT"]]
tables["Tab6_Cost"] = (co, "Table VI: training time (minutes, 30 epochs, RTX 3090) and overhead relative to PGD-AT.", [("bar_cost", None)])

# ============================================================ ALL results (every run, every epoch, every class)
import sys, re, zipfile
sys.path.insert(0, HERE)
from common import load_dataset  # noqa: E402

DSN = {"cifar10": "CIFAR-10", "cifar100": "CIFAR-100", "fashionmnist": "Fashion-MNIST", "tinyimagenet": "Tiny-ImageNet (15 cl.)"}
runs = sorted(os.path.relpath(os.path.dirname(p), RES).replace("\\", "/") for p in glob.glob(os.path.join(RES, "*", "*", "*", "DONE")))
EPS_KEYS = ["0", "1", "2", "4", "6", "8", "10", "12", "16"]
STEP_KEYS = ["1", "2", "5", "10", "20", "50", "100"]
all_runs, all_epochs, all_pc = [], [], []
labels = {}
for run in runs:
    ds, bb, folder = run.split("/")
    h = hi(ds, bb, folder); e = ev(ds, bb, folder); a = h["args"]
    r = {"Dataset": DSN[ds], "Backbone": BN.get(bb, bb), "Run": folder, "Method": a["method"],
         "Params (M)": round(h["n_params"] / 1e6, 3), "Epochs": a["epochs"], "Batch": a["batch_size"], "LR": a["lr"],
         "Weight decay": a["wd"], "Warmup": a["warmup"], "Seed": a["seed"],
         "EMA alpha": a.get("ema_alpha") if a["method"] == "acat" else None,
         "beta (adv share)": a.get("beta_adv", 0.5) if a["method"] == "acat" else None,
         "lambda mode": a.get("lam_mode", "sched") if a["method"] == "acat" else None,
         "schedule": a.get("schedule", "linear") if a["method"] == "acat" else None,
         "h direction": a.get("h_dir", "easy") if a["method"] == "acat" else None,
         "Train time (min)": round(h["train_time_sec"] / 60, 2),
         "Evaluated": e is not None and bool(e.get("done"))}
    for k, n in ATT:
        r[n] = acc(e, k)
    r["Transfer source"] = e.get("transfer_source") if e else None
    for k in EPS_KEYS:
        r[f"PGD-20 @ eps={k}/255"] = round(100 * e["eps_curve"][k], 2) if e and "eps_curve" in e else None
    for k in STEP_KEYS:
        r[f"PGD-{k} @ 8/255"] = round(100 * e["steps_curve"][k], 2) if e and "steps_curve" in e else None
    if e:
        r["Eval time (min)"] = round(sum(e.get("eval_time_sec", {}).values()) / 60, 2)
    all_runs.append(r)
    for ep in h["epochs"]:
        row = {"Dataset": DSN[ds], "Backbone": BN.get(bb, bb), "Run": folder, "Epoch": ep["epoch"] + 1}
        row.update({k: (round(v, 5) if isinstance(v, float) else v) for k, v in ep.items() if k != "epoch"})
        all_epochs.append(row)
    if e and "preds" in e.get("pgd20", {}):
        if ds not in labels:
            labels[ds] = load_dataset(ds)[3]
        yy = labels[ds]
        cp, pp = np.array(e["clean"]["preds"]), np.array(e["pgd20"]["preds"])
        for c in range(int(yy.max()) + 1):
            m_ = yy == c
            all_pc.append({"Dataset": DSN[ds], "Backbone": BN.get(bb, bb), "Run": folder, "Class": c,
                           "Class name": CLS[c] if ds == "cifar10" else "", "n test": int(m_.sum()),
                           "Clean acc": round(100 * float((cp[m_] == c).mean()), 2),
                           "PGD-20 acc": round(100 * float((pp[m_] == c).mean()), 2)})

tables["All_Runs"] = (pd.DataFrame(all_runs), "ALL GPU runs (every dataset, backbone, method and ablation): settings, accuracy (%) under every attack, eps and PGD-step curves, train/eval time. Runs with Evaluated=False are transfer-attack source models only.", [])
tables["All_Epochs"] = (pd.DataFrame(all_epochs), "ALL GPU runs, every epoch: training loss, epoch time, monitor clean / PGD-10 accuracy (fractions, first 1,000 test images) and, for ACAT runs, curriculum statistics.", [])
tables["All_PerClass"] = (pd.DataFrame(all_pc), "ALL evaluated GPU runs: per-class clean and PGD-20 accuracy (%) on the full test set.", [])

# ---- earlier studies (kept for completeness; NOT used in the v7 paper)
ER = ROOT
earlier = [
    ("Earlier_CPU_v2_summary", "results/results_summary_v2.csv",
     "EARLIER CPU study (paper v5/v6 Table V): SmallCNN, 640/800/450 training images, 20/18/10 epochs, 150/150/100 test images, 2 seeds (1 for Tiny). Accuracies in %, MES relative to 5 methods. Superseded by the GPU study."),
    ("Earlier_CPU_v2_per_seed", "results/results_summary_v2_per_seed.csv",
     "EARLIER CPU study, per seed (accuracies as fractions). Superseded."),
    ("Earlier_CPU_v1_4to6ep", "results_v1_singlerun_backup/results_summary.csv",
     "EARLIER preliminary CPU study, 4-6 epochs, single seed (accuracies as fractions). Superseded."),
    ("Earlier_CPU_v5_backbones", "results_full_matrix/comparison_summary_full.csv",
     "EARLIER CPU smoke test, 5 backbones, 64 training / 24 test images, CIFAR-10. Near chance level; superseded."),
    ("Earlier_CPU_v6_matrix", "results_gpu_matrix_quick/comparison_summary_full_v6.csv",
     "EARLIER CPU smoke test, 3 datasets x 4 backbones, 4 epochs, 64 training / 24 test images. Near chance level; superseded."),
    ("Earlier_CPU_v4_backbones", "results_multi_model/comparison_summary.csv",
     "EARLIER CPU smoke test (paper v5 Table VII), 64 training / 40 test images. Near chance level; superseded."),
]
for name, rel, desc in earlier:
    p = os.path.join(ER, rel)
    if os.path.exists(p):
        tables[name] = (pd.read_csv(p), desc, [])

# ---- tables of the v5 manuscript (parsed from its LaTeX source, not re-typed)
z = os.path.join(os.path.dirname(ROOT), "old", "ACAT_paper_latex_source_v5.zip")
if os.path.exists(z):
    tex = zipfile.ZipFile(z).read("paper_clean/main_ieee.tex").decode("utf8")

    def parse_table(label):
        i = tex.index(r"\label{" + label + "}")
        j = tex.index(r"\end{tabular}", i)
        body = tex[i:j].replace(r"\&", "AMPERSAND")  # protect C\&W before splitting cells on &
        head = body[body.index(r"\toprule") + len(r"\toprule"): body.index(r"\midrule")]
        rows_ = body[body.index(r"\midrule") + len(r"\midrule"): body.index(r"\bottomrule")]
        clean = lambda t: re.sub(r"\\cite\{[^}]*\}|\\textbf\{([^}]*)\}|\\multirow\{\d+\}\{\*\}\{([^}]*)\}|\\,s|\\%|\\&|\$|~", lambda m: (m.group(1) or m.group(2) or ("&" if m.group(0) == r"\&" else "")), t).replace(r"\pm", "+/-").replace("AMPERSAND", "&").replace(r"\%", "%").replace("$", "").strip()
        cols = [clean(c) for c in head.replace(r"\\", "").split("&")]
        out = []
        for line in rows_.split(r"\\"):
            line = line.replace(r"\midrule", "").strip()
            if not line:
                continue
            out.append([clean(c) for c in line.split("&")])
        df_ = pd.DataFrame(out, columns=cols[:len(out[0])] if len(cols) >= len(out[0]) else None)
        first = df_.columns[0]
        df_[first] = df_[first].replace("", np.nan).ffill()  # fill \multirow continuation rows
        return df_

    for label, name, desc in [
        ("tab:cifar10", "v5_Paper_CIFAR10_WRN", "v5 MANUSCRIPT Table II (CIFAR-10, WRN-28-10, 150 epochs) from a separate, unreleased codebase; NOT reproduced by the GPU study and removed from the v7 paper."),
        ("tab:cifar100", "v5_Paper_CIFAR100_WRN", "v5 MANUSCRIPT Table III (CIFAR-100, WRN-28-10); same caveat as above."),
        ("tab:ablation", "v5_Paper_Ablation_WRN", "v5 MANUSCRIPT Table IV (ablation, CIFAR-10, WRN-28-10); same caveat as above."),
        ("tab:extended", "v5_Paper_CPU_Table", "v5 MANUSCRIPT Table V (CPU study) as printed; note the Tiny-ImageNet ACAT clean value 27.00 was a typo (CSV: 26.00)."),
        ("tab:wrn_feasibility", "v5_Paper_WRN_Smoke", "v5 MANUSCRIPT Table VI (WRN-28-10 CPU smoke test, 64 training / 40 test images)."),
    ]:
        try:
            tables[name] = (parse_table(label), desc, [])
        except Exception as ex:  # keep going if one table has an unusual layout
            print("could not parse", label, ex)

# ============================================================ write CSVs
for name, (df, _, _) in tables.items():
    df.to_csv(os.path.join(CSV, f"{name}.csv"), index=False)

# ============================================================ write Excel
wb = Workbook()
ws0 = wb.active; ws0.title = "README"
ws0["A1"] = "ACAT paper - data behind every figure and table"; ws0["A1"].font = Font(bold=True, size=14)
info = ["Source: acat_gpu/results_gpu (RTX 3090, 30 epochs, seed 42, full datasets).",
        "Attacks: L-inf eps=8/255; FGSM, PGD-20, CW-Linf (30 steps), AutoAttack standard (first 1,000 test images), Square (1,000 queries, first 1,000 images), transfer from a standard model of a different architecture.",
        "MES is computed among the four AT methods per setting (Standard excluded).",
        "Figs. 1-3 of the paper are illustrations (concept, pipeline, example images) and have no underlying data table.", ""]
for i, t in enumerate(info, start=2):
    ws0.cell(row=i, column=1, value=t)
ws0.cell(row=8, column=1, value="Sheet").font = Font(bold=True)
ws0.cell(row=8, column=2, value="Contents").font = Font(bold=True)
for i, (name, (_, desc, _)) in enumerate(tables.items(), start=9):
    c = ws0.cell(row=i, column=1, value=name); c.hyperlink = f"#'{name}'!A1"; c.font = Font(color="0563C1", underline="single")
    ws0.cell(row=i, column=2, value=desc)
ws0.column_dimensions["A"].width = 28; ws0.column_dimensions["B"].width = 140

HEAD = PatternFill("solid", fgColor="DDEBF7")


def write_df(ws, df, r0=3):
    for j, col in enumerate(df.columns, start=1):
        c = ws.cell(row=r0, column=j, value=col); c.font = Font(bold=True); c.fill = HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(j)].width = max(11, min(26, len(str(col)) + 3))
    for i, row in enumerate(df.itertuples(index=False), start=r0 + 1):
        for j, v in enumerate(row, start=1):
            ws.cell(row=i, column=j, value=None if (isinstance(v, float) and np.isnan(v)) else v)
    ws.freeze_panes = ws.cell(row=r0 + 1, column=1)
    return r0


def color_series(chart):
    for s in chart.series:
        name = s.tx.strRef.f if s.tx and s.tx.strRef else ""
        for k, col in COLORS.items():
            if s.tx and s.tx.strRef and k in str(s.tx.v if hasattr(s.tx, "v") else ""):
                pass
    return chart


def series_color(s, hexcol, line=True):
    if line:
        s.graphicalProperties.line.solidFill = hexcol
        s.graphicalProperties.line.width = 22000
        s.smooth = False
    else:
        s.graphicalProperties.solidFill = hexcol
        s.graphicalProperties.line.solidFill = hexcol


def line_chart(ws, title, xtitle, ytitle, xcol, ycols, r1, r2, r0, anchor, names=None, w=16, h=8):
    ch = LineChart(); ch.title = title; ch.x_axis.title = xtitle; ch.y_axis.title = ytitle
    ch.width, ch.height = w, h
    for yc in ycols:
        ref = Reference(ws, min_col=yc, min_row=r1, max_row=r2)
        s = Series(ref, title=names[yc] if names else ws.cell(row=r0, column=yc).value)
        s.smooth = False
        ch.series.append(s)
        nm = names[yc] if names else ws.cell(row=r0, column=yc).value
        for k, col in COLORS.items():
            if str(nm).startswith(k):
                series_color(s, col)
    ch.set_categories(Reference(ws, min_col=xcol, min_row=r1, max_row=r2))
    ch.x_axis.delete = False; ch.y_axis.delete = False
    ws.add_chart(ch, anchor)


def bar_chart(ws, title, ytitle, catcol, ycols, r1, r2, r0, anchor, w=18, h=8):
    ch = BarChart(); ch.type = "col"; ch.title = title; ch.y_axis.title = ytitle
    ch.width, ch.height = w, h
    for yc in ycols:
        s = Series(Reference(ws, min_col=yc, min_row=r1, max_row=r2), title=ws.cell(row=r0, column=yc).value)
        nm = str(ws.cell(row=r0, column=yc).value)
        for k, col in COLORS.items():
            if nm.startswith(k):
                series_color(s, col, line=False)
        ch.series.append(s)
    ch.set_categories(Reference(ws, min_col=catcol, min_row=r1, max_row=r2))
    ch.x_axis.delete = False; ch.y_axis.delete = False
    ws.add_chart(ch, anchor)


for name, (df, desc, charts) in tables.items():
    ws = wb.create_sheet(name[:31])
    ws["A1"] = desc; ws["A1"].font = Font(italic=True)
    r0 = write_df(ws, df)
    n = len(df); cols = list(df.columns)
    chart_col = get_column_letter(len(cols) + 2)
    for spec in charts:
        kind = spec[0]
        if kind == "bar":
            bar_chart(ws, name.split("_", 1)[1], spec[1], 1, list(range(2, len(cols) + 1)), r0 + 1, r0 + n, r0, f"{chart_col}3", w=24)
        elif kind == "bar_blocks":
            groups = df.groupby(spec[1], sort=False)
            row = 3
            start = r0 + 1
            for g, gdf in groups:
                ycols = [cols.index(c) + 1 for c in cols if c in COLORS]
                bar_chart(ws, g, "Accuracy (%)", 2, ycols, start, start + len(gdf) - 1, r0, f"{chart_col}{row}", w=22, h=7.5)
                start += len(gdf); row += 16
        elif kind == "line_blocks":
            _, gcol, xname = spec
            start = r0 + 1; row = 3
            for g, gdf in df.groupby(gcol, sort=False):
                ycols = [cols.index(c) + 1 for c in cols if c in COLORS]
                line_chart(ws, f"CIFAR-10, {g}", xname, "Robust accuracy (%)", cols.index(xname) + 1, ycols,
                           start, start + len(gdf) - 1, r0, f"{chart_col}{row}")
                start += len(gdf); row += 17
        elif kind in ("line_dyn", "line_dyn_ds"):
            gcol = "Backbone" if kind == "line_dyn" else "Dataset"
            start = r0 + 1; row = 3
            for g, gdf in df.groupby(gcol, sort=False):
                for k, (metric, yl) in enumerate([("loss", "Training loss"), ("clean acc", "Clean accuracy (%)"), ("PGD-10 acc", "PGD-10 accuracy (%)")]):
                    ycols = [cols.index(f"{MN[m]} {metric}") + 1 for m in METHODS]
                    nm = {c: cols[c - 1].replace(f" {metric}", "") for c in ycols}
                    line_chart(ws, f"{g}: {yl}", "Epoch", yl, cols.index("Epoch") + 1, ycols, start, start + len(gdf) - 1, r0,
                               f"{get_column_letter(len(cols) + 2 + 10 * k)}{row}", names=nm, w=15, h=7.5)
                start += len(gdf); row += 16
        elif kind == "line_cur":
            start = r0 + 1; row = 3
            for g, gdf in df.groupby("Backbone", sort=False):
                e1 = start + len(gdf) - 1
                line_chart(ws, f"{g}: correlation of weights", "Epoch", "Pearson r", 2,
                           [cols.index(c) + 1 for c in ["lambda_t", "corr(w,h)", "corr(w,b_smooth)"]], start, e1, r0, f"{chart_col}{row}", w=15, h=7.5)
                line_chart(ws, f"{g}: spread of weights", "Epoch", "w", 2,
                           [cols.index(c) + 1 for c in ["std(w)", "w 10th pct", "w 90th pct"]], start, e1, r0,
                           f"{get_column_letter(len(cols) + 12)}{row}", w=15, h=7.5)
                line_chart(ws, f"{g}: mean signals", "Epoch", "value", 2,
                           [cols.index(c) + 1 for c in ["mean h", "mean b_smooth", "mean b (raw)"]], start, e1, r0,
                           f"{get_column_letter(len(cols) + 22)}{row}", w=15, h=7.5)
                start += len(gdf); row += 16
        elif kind == "line_bins":
            ycols = [j + 1 for j, c in enumerate(cols) if c.startswith("mean w")]
            line_chart(ws, "Mean curriculum weight vs. hardness", "hardness h (bin center)", "mean w", 1, ycols, r0 + 1, r0 + n, r0, f"{chart_col}3", w=20, h=10)
        elif kind == "bar_perclass":
            row = 3
            sub = [(i, r) for i, r in enumerate(df.itertuples(index=False)) if r.Metric == "PGD-20 acc"]
            # helper table: transposed PGD-20 per class for each backbone
            hc = len(cols) + 2
            rr = r0
            for bb in BN.values():
                part = df[(df["Backbone"] == bb) & (df["Metric"] == "PGD-20 acc")]
                ws.cell(row=rr, column=hc, value=f"{bb} class").font = Font(bold=True)
                for j, m in enumerate(part["Method"], start=1):
                    ws.cell(row=rr, column=hc + j, value=m).font = Font(bold=True)
                for i, c in enumerate(CLS, start=1):
                    ws.cell(row=rr + i, column=hc, value=c)
                    for j, (_, prow) in enumerate(part.iterrows(), start=1):
                        ws.cell(row=rr + i, column=hc + j, value=prow[c])
                bar_chart(ws, f"Per-class PGD-20 accuracy, {bb}", "Accuracy (%)", hc, list(range(hc + 1, hc + 6)), rr + 1, rr + 10, rr,
                          f"{get_column_letter(hc + 7)}{rr}", w=20, h=7.5)
                rr += 16
        elif kind == "bar_cost":
            bar_chart(ws, "Training time (min)", "minutes", 1, [2, 3, 4, 5], r0 + 1, r0 + n, r0, f"{chart_col}3", w=20)
        elif kind == "scatter_methods":
            ch = ScatterChart(); ch.title = "CIFAR-10: clean vs AutoAttack"; ch.style = 13
            ch.x_axis.title = "Clean accuracy (%)"; ch.y_axis.title = "AutoAttack accuracy (%)"; ch.width, ch.height = 18, 10
            for m in ["PGD-AT", "TRADES", "MART", "ACAT"]:
                idx = [i for i, v in enumerate(df["Method"]) if v == m]
                # contiguous helper rows per method
                hc = len(cols) + 2 + ["PGD-AT", "TRADES", "MART", "ACAT"].index(m) * 3
                ws.cell(row=r0, column=hc, value=f"{m} clean").font = Font(bold=True)
                ws.cell(row=r0, column=hc + 1, value=f"{m} AA").font = Font(bold=True)
                for k, i in enumerate(idx, start=1):
                    ws.cell(row=r0 + k, column=hc, value=df["Clean"].iloc[i]); ws.cell(row=r0 + k, column=hc + 1, value=df["AutoAttack"].iloc[i])
                s = Series(Reference(ws, min_col=hc + 1, min_row=r0 + 1, max_row=r0 + len(idx)),
                           Reference(ws, min_col=hc, min_row=r0 + 1, max_row=r0 + len(idx)), title=m)
                s.marker.symbol = "circle"; s.marker.size = 9; s.graphicalProperties.line.noFill = True
                s.marker.graphicalProperties.solidFill = COLORS[m]; s.marker.graphicalProperties.line.solidFill = COLORS[m]
                ch.series.append(s)
            ch.x_axis.scaling.min = 70; ch.x_axis.scaling.max = 90
            ch.y_axis.scaling.min = 25; ch.y_axis.scaling.max = 55
            ch.x_axis.delete = False; ch.y_axis.delete = False
            ws.add_chart(ch, f"{get_column_letter(len(cols) + 15)}3")
        elif kind == "scatter_ablation":
            ch = ScatterChart(); ch.title = "Ablation: clean vs AutoAttack (ResNet-18)"; ch.style = 13
            ch.x_axis.title = "Clean accuracy (%)"; ch.y_axis.title = "AutoAttack accuracy (%)"; ch.width, ch.height = 20, 11
            ci, ai = cols.index("Clean") + 1, cols.index("AutoAttack") + 1
            for i in range(n):
                s = Series(Reference(ws, min_col=ai, min_row=r0 + 1 + i), Reference(ws, min_col=ci, min_row=r0 + 1 + i),
                           title=df["Variant"].iloc[i])
                s.marker.symbol = "diamond" if df["folder"].iloc[i].startswith("acat") else "circle"
                s.marker.size = 10; s.graphicalProperties.line.noFill = True
                f = df["folder"].iloc[i]
                col = {"acat": COLORS["ACAT"], "pgd_at": COLORS["PGD-AT"], "trades": COLORS["TRADES"], "mart": COLORS["MART"]}.get(
                    f, "6B6B6B" if "adv_only" in f else "E08A2C")
                s.marker.graphicalProperties.solidFill = col; s.marker.graphicalProperties.line.solidFill = col
                ch.series.append(s)
            ch.x_axis.scaling.min = 76; ch.x_axis.scaling.max = 88
            ch.y_axis.scaling.min = 30; ch.y_axis.scaling.max = 50
            ch.x_axis.delete = False; ch.y_axis.delete = False
            ws.add_chart(ch, f"{chart_col}3")

wb.save(XLSX)
print("saved", XLSX)
print("csv:", sorted(os.listdir(CSV)))
